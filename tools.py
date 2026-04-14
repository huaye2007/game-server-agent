"""工具定义 - 给AI提供文件搜索、读取和写入能力"""
import os
import re
from pathlib import Path
from i18n import t


def list_dir(project_path: str, rel_dir: str = "", max_depth: int = 2,
             ignore_dirs: list = None, **_kwargs) -> str:
    """列出目录结构"""
    root = Path(project_path) / rel_dir
    if not root.exists():
        return t("dir_not_exist", path=rel_dir)
    ignore = set(ignore_dirs or [])
    lines = []
    _walk_tree(root, Path(project_path), lines, ignore, 0, max_depth)
    if not lines:
        return t("empty_dir")
    return "\n".join(lines[:200])


def _walk_tree(path: Path, root: Path, lines: list, ignore: set,
               depth: int, max_depth: int):
    if depth > max_depth:
        return
    try:
        entries = sorted(path.iterdir(), key=lambda p: (p.is_file(), p.name))
    except PermissionError:
        return
    for entry in entries:
        if entry.name.startswith(".") or entry.name in ignore:
            continue
        rel = entry.relative_to(root)
        prefix = "  " * depth
        if entry.is_dir():
            lines.append(f"{prefix}{rel}/")
            _walk_tree(entry, root, lines, ignore, depth + 1, max_depth)
        else:
            lines.append(f"{prefix}{rel}")


def read_file(project_path: str, rel_path: str, max_lines: int = 500,
              **_kwargs) -> str:
    """读取文件内容"""
    fp = Path(project_path) / rel_path
    if not fp.exists():
        return t("file_not_exist", path=rel_path)
    if not fp.is_file():
        return t("not_a_file", path=rel_path)
    try:
        with open(fp, "r", encoding="utf-8", errors="ignore") as f:
            lines = f.readlines()
        if len(lines) > max_lines:
            half = max_lines // 2
            head = "".join(lines[:half])
            tail = "".join(lines[-half:])
            return f"{head}\n{t('lines_omitted', count=len(lines) - max_lines)}\n{tail}"
        return "".join(lines)
    except Exception as e:
        return t("read_error", err=e)


def search_files(project_path: str, pattern: str,
                 ignore_dirs: list = None, **_kwargs) -> str:
    """按文件名模糊搜索"""
    root = Path(project_path)
    ignore = set(ignore_dirs or [])
    results = []
    pat = pattern.lower()
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames
                       if d not in ignore and not d.startswith(".")]
        for fname in filenames:
            if pat in fname.lower():
                rel = str(Path(dirpath, fname).relative_to(root))
                results.append(rel)
                if len(results) >= 50:
                    return "\n".join(results) + "\n" + t("results_truncated")
    return "\n".join(results) if results else t("no_match_files")


def grep_search(project_path: str, pattern: str, file_ext: str = "",
                ignore_dirs: list = None, max_results: int = 30,
                **_kwargs) -> str:
    """在文件内容中搜索文本"""
    root = Path(project_path)
    ignore = set(ignore_dirs or [])
    results = []
    try:
        regex = re.compile(pattern, re.IGNORECASE)
    except re.error:
        regex = re.compile(re.escape(pattern), re.IGNORECASE)

    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames
                       if d not in ignore and not d.startswith(".")]
        for fname in filenames:
            if file_ext and not fname.endswith(file_ext):
                continue
            fp = Path(dirpath, fname)
            try:
                size = fp.stat().st_size
                if size > 512 * 1024:
                    continue
                with open(fp, "r", encoding="utf-8", errors="ignore") as f:
                    for i, line in enumerate(f, 1):
                        if regex.search(line):
                            rel = str(fp.relative_to(root))
                            results.append(f"{rel}:{i}: {line.rstrip()[:150]}")
                            if len(results) >= max_results:
                                return "\n".join(results) + "\n" + t("results_truncated")
            except Exception:
                continue
    return "\n".join(results) if results else t("no_match_content")


# ============ 写文件工具 ============

def write_file(project_path: str, rel_path: str, content: str,
               base_dir: str = "", **_kwargs) -> str:
    """写入文件内容到指定路径。"""
    if base_dir:
        fp = Path(base_dir) / rel_path
    else:
        fp = Path(project_path) / rel_path
    try:
        fp.parent.mkdir(parents=True, exist_ok=True)
        with open(fp, "w", encoding="utf-8") as f:
            f.write(content)
        return t("file_written", path=fp)
    except Exception as e:
        return t("write_error", err=e)


# ============ 协议文件目录工具 ============

def list_proto_dir(project_path: str, proto_path: str = "",
                   rel_dir: str = "", max_depth: int = 2,
                   **_kwargs) -> str:
    """列出协议文件(.proto)目录结构"""
    if not proto_path:
        return t("proto_not_set")
    root = Path(proto_path) / rel_dir
    if not root.exists():
        return t("proto_dir_not_exist", path=proto_path)
    lines = []
    _walk_tree(root, Path(proto_path), lines, set(), 0, max_depth)
    if not lines:
        return t("empty_dir")
    return "\n".join(lines[:200])


def read_proto_file(project_path: str, proto_path: str = "",
                    rel_path: str = "", **_kwargs) -> str:
    """读取协议文件(.proto)内容"""
    if not proto_path:
        return t("proto_not_set_short")
    fp = Path(proto_path) / rel_path
    if not fp.exists():
        return t("file_not_exist", path=rel_path)
    try:
        return fp.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        return t("read_error", err=e)


# ============ 配置表目录工具 ============

def list_excel_dir(project_path: str, excel_path: str = "",
                   rel_dir: str = "", max_depth: int = 2,
                   **_kwargs) -> str:
    """列出配置表(Excel)目录结构"""
    if not excel_path:
        return t("excel_not_set")
    root = Path(excel_path) / rel_dir
    if not root.exists():
        return t("excel_dir_not_exist", path=excel_path)
    lines = []
    _walk_tree(root, Path(excel_path), lines, set(), 0, max_depth)
    if not lines:
        return t("empty_dir")
    return "\n".join(lines[:200])


def read_excel_meta(project_path: str, excel_path: str = "",
                    rel_path: str = "", **_kwargs) -> str:
    """读取Excel配置表的sheet名和表头信息"""
    if not excel_path:
        return t("excel_not_set_short")
    fp = Path(excel_path) / rel_path
    if not fp.exists():
        return t("file_not_exist", path=rel_path)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
        result = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(max_row=5, values_only=True))
            result.append(f"## Sheet: {sheet_name}")
            for i, row in enumerate(rows):
                cells = [str(c) if c is not None else "" for c in row]
                result.append(f"  行{i+1}: {' | '.join(cells)}")
        wb.close()
        return "\n".join(result)
    except ImportError:
        return t("need_openpyxl")
    except Exception as e:
        return t("read_error", err=e)


def write_excel(project_path: str, excel_path: str = "",
                rel_path: str = "", sheets: str = "", **_kwargs) -> str:
    """写入Excel配置表。"""
    import json as _json
    if not excel_path:
        return t("excel_not_set_short")
    fp = Path(excel_path) / rel_path
    try:
        import openpyxl
        sheet_data = _json.loads(sheets)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sd in sheet_data:
            ws = wb.create_sheet(title=sd["name"])
            if sd.get("headers"):
                ws.append(sd["headers"])
            for row in sd.get("rows", []):
                ws.append(row)
        fp.parent.mkdir(parents=True, exist_ok=True)
        wb.save(fp)
        return t("excel_written", path=fp)
    except ImportError:
        return t("need_openpyxl")
    except Exception as e:
        return t("write_error", err=e)


# ============ 记忆文件（AGENT.md） ============

MEMORY_FILE = "AGENT.md"


def memory_read(project_path: str, **_kwargs) -> str:
    """读取项目记忆文件 AGENT.md"""
    fp = Path(project_path) / MEMORY_FILE
    if not fp.exists():
        return t("memory_not_exist")
    try:
        return fp.read_text(encoding="utf-8")
    except Exception as e:
        return t("read_error", err=e)


def memory_append(project_path: str, content: str, **_kwargs) -> str:
    """向项目记忆文件 AGENT.md 追加内容"""
    fp = Path(project_path) / MEMORY_FILE
    try:
        existing = ""
        if fp.exists():
            existing = fp.read_text(encoding="utf-8")
        with open(fp, "w", encoding="utf-8") as f:
            if existing and not existing.endswith("\n"):
                existing += "\n"
            f.write(existing + content.strip() + "\n")
        return t("memory_written")
    except Exception as e:
        return t("write_error", err=e)


def load_memory(project_path: str) -> str:
    """启动时加载记忆文件内容"""
    fp = Path(project_path) / MEMORY_FILE
    if not fp.exists():
        return ""
    try:
        return fp.read_text(encoding="utf-8")
    except Exception:
        return ""


# ============ 工具定义（OpenAI function calling 格式） ============

TOOL_DEFINITIONS = [
    {
        "type": "function",
        "function": {
            "name": "list_dir",
            "description": "List project directory structure. Use to understand project layout and find relevant directories.",
            "parameters": {
                "type": "object",
                "properties": {
                    "rel_dir": {
                        "type": "string",
                        "description": "Subdirectory path relative to project root, empty string for root"
                    },
                    "max_depth": {
                        "type": "integer",
                        "description": "Directory depth, default 2"
                    }
                },
                "required": []
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "read_file",
            "description": "Read file content from the project. Use to view code implementation details.",
            "parameters": {
                "type": "object",
                "properties": {
                    "rel_path": {
                        "type": "string",
                        "description": "File path relative to project root"
                    }
                },
                "required": ["rel_path"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "search_files",
            "description": "Fuzzy search files by name in the project. Use to find specific file types or names.",
            "parameters": {
                "type": "object",
                "properties": {
                    "pattern": {
                        "type": "string",
                        "description": "Filename search keyword (fuzzy match)"
                    }
                },
                "required": ["pattern"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "grep_search",
            "description": "Search text or regex in project file contents. Use to find class names, method names, protocol definitions, etc.",
            "parameters": {
                "type": "object",
                "properties": {
                    "pattern": {
                        "type": "string",
                        "description": "Text or regex pattern to search"
                    },
                    "file_ext": {
                        "type": "string",
                        "description": "Limit to file extension, e.g. .java .go .lua, empty for all"
                    }
                },
                "required": ["pattern"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "write_file",
            "description": "Write file to project directory. Use to generate business code, protocol files, etc. Only call after user confirmation.",
            "parameters": {
                "type": "object",
                "properties": {
                    "rel_path": {
                        "type": "string",
                        "description": "Relative file path"
                    },
                    "content": {
                        "type": "string",
                        "description": "File content"
                    },
                    "base_dir": {
                        "type": "string",
                        "description": "Base directory: 'proto' for proto dir, 'excel' for config dir, empty for project dir"
                    }
                },
                "required": ["rel_path", "content"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "list_proto_dir",
            "description": "List .proto file directory structure to view existing protocol files.",
            "parameters": {
                "type": "object",
                "properties": {
                    "rel_dir": {
                        "type": "string",
                        "description": "Subdirectory path relative to proto directory, empty for root"
                    },
                    "max_depth": {
                        "type": "integer",
                        "description": "Directory depth, default 2"
                    }
                },
                "required": []
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "read_proto_file",
            "description": "Read existing .proto file content for reference on protocol format and style.",
            "parameters": {
                "type": "object",
                "properties": {
                    "rel_path": {
                        "type": "string",
                        "description": "Proto file path relative to proto directory"
                    }
                },
                "required": ["rel_path"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "list_excel_dir",
            "description": "List Excel config table directory structure to view existing config files.",
            "parameters": {
                "type": "object",
                "properties": {
                    "rel_dir": {
                        "type": "string",
                        "description": "Subdirectory path relative to Excel directory, empty for root"
                    },
                    "max_depth": {
                        "type": "integer",
                        "description": "Directory depth, default 2"
                    }
                },
                "required": []
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "read_excel_meta",
            "description": "Read Excel config table sheet names and first few rows for reference on existing table format and field design.",
            "parameters": {
                "type": "object",
                "properties": {
                    "rel_path": {
                        "type": "string",
                        "description": "Excel file path relative to Excel directory"
                    }
                },
                "required": ["rel_path"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "write_excel",
            "description": "Write Excel config table file. Only call after user confirmation.",
            "parameters": {
                "type": "object",
                "properties": {
                    "rel_path": {
                        "type": "string",
                        "description": "Excel file path relative to Excel directory"
                    },
                    "sheets": {
                        "type": "string",
                        "description": "JSON string, format: [{\"name\": \"Sheet1\", \"headers\": [\"col1\",\"col2\"], \"rows\": [[\"val1\",\"val2\"]]}]"
                    }
                },
                "required": ["rel_path", "sheets"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "memory_read",
            "description": "Read project memory file AGENT.md. Contains project conventions, lessons learned, and development rules.",
            "parameters": {
                "type": "object",
                "properties": {},
                "required": []
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "memory_append",
            "description": "Append content to project memory file AGENT.md. ONLY for universal project rules and conventions (naming, code style, framework patterns). Do NOT record module-specific details or individual feature notes.",
            "parameters": {
                "type": "object",
                "properties": {
                    "content": {
                        "type": "string",
                        "description": "Memory content to append (markdown format)"
                    }
                },
                "required": ["content"]
            }
        }
    },
]
